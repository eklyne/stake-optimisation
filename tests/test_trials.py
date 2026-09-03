"""Tests for `sims.pptx` - the one-lifetime-per-slide deck.

Three things are worth guarding here. The first is the ARITHMETIC on the corner
table: every row is measured off the path rather than sampled, so a hand-built
path with a known worst fall pins all three drawdown rows at once. The second is
the SHARED AXES - the deck exists to be flicked through, and a chart that
autoscaled per slide would quietly undo the whole point of it.

The third is the RAKEBACK SPLIT. The slides draw the lifetime as three lines and
price the fall twice, so the guarantee that matters is that the split is exact:
the parts sum to the total at every step and in every row, the fall is the fall
in what the cards paid, and the net figure is identically the distance the
bankroll moved over the same stretch rather than a second measurement of it.
"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import numpy as np

from shotopt import mix, trials
from shotopt.config import Config, Stake


def _config(**overrides) -> Config:
    defaults = dict(
        bankroll_eur=40_000.0,
        tables=12,
        ruin_tolerance=0.01,
        timescale_hands=200_000,
        stakes=(
            Stake("100NL", 1.0, 7.46, 92.0),
            Stake("200NL", 2.0, 4.32, 92.0),
        ),
        rakeback_pct=0.0,
    )
    defaults.update(overrides)
    return Config(**defaults)


class TestTrialStatistics(unittest.TestCase):
    """The corner table, against a path whose answers are known by hand."""

    def setUp(self):
        self.config = _config()
        self.allocation = mix.evaluate((4, 8), self.config)

    def _describe(self, equity, rakeback=None):
        equity = np.asarray(equity, dtype=float)
        rakeback = (np.zeros_like(equity) if rakeback is None
                    else np.asarray(rakeback, dtype=float))
        hands = np.arange(len(equity)) * trials._HANDS_PER_STEP
        return trials._describe(self.config, self.allocation, equity, rakeback,
                                hands, False)

    def test_the_worst_fall_is_the_deepest_one_not_the_last(self):
        # Up to 100, down to 40 (a fall of 60), back to 120, down to 90 (a fall
        # of 30). The measured fall must be the 60, and its length the two steps
        # the 60 took - NOT the shallower fall that happens to end the path.
        stats = self._describe([0, 100, 70, 40, 120, 90])
        self.assertAlmostEqual(stats["drawdown_eur"], 60.0)
        self.assertEqual(stats["drawdown_from"], 1 * trials._HANDS_PER_STEP)
        self.assertEqual(stats["drawdown_to"], 3 * trials._HANDS_PER_STEP)
        self.assertEqual(stats["drawdown_hands"], 2 * trials._HANDS_PER_STEP)

    def test_the_fall_is_timed_from_the_first_step_at_the_high(self):
        # A flat plateau at the top is not part of the fall. Timing it from the
        # LAST step at the high would be the same depth over a shorter stretch,
        # which is the reading that flatters the mix.
        stats = self._describe([0, 100, 100, 100, 40])
        self.assertAlmostEqual(stats["drawdown_eur"], 60.0)
        self.assertEqual(stats["drawdown_hands"], 3 * trials._HANDS_PER_STEP)

    def test_recovery_is_measured_from_the_trough_back_to_the_old_high(self):
        stats = self._describe([0, 100, 40, 60, 100, 130])
        self.assertEqual(stats["recovery_hands"], 2 * trials._HANDS_PER_STEP)

    def test_a_fall_never_recovered_reads_as_none_not_as_zero(self):
        # It is quoted on the slide as "never recovered". Zero would read as
        # "recovered instantly", which is the opposite of what happened.
        stats = self._describe([0, 100, 40, 60, 70])
        self.assertIsNone(stats["recovery_hands"])

    def test_the_high_water_mark_includes_the_start(self):
        # A lifetime that never gets above water still has its fall measured
        # from zero, so a losing path cannot report a small drawdown just
        # because it was never up.
        stats = self._describe([0, -30, -80, -50])
        self.assertAlmostEqual(stats["drawdown_eur"], 80.0)
        self.assertAlmostEqual(stats["worst_below_start_eur"], 80.0)

    def test_buyins_and_bb100_agree_with_the_money(self):
        stats = self._describe([0, 100, 2160])
        buyin = trials._buyin_eur(self.allocation)
        # 4x 100NL + 8x 200NL: a 100bb buy-in averages EUR 166.67 across them.
        self.assertAlmostEqual(buyin, (4 * 100.0 + 8 * 200.0) / 12)
        self.assertAlmostEqual(stats["won_buyins"], 2160.0 / buyin)
        # bb/100 is the same money on the same converter: bb won over hundreds
        # of hands. It must reconcile exactly, not approximately.
        self.assertAlmostEqual(
            stats["bb100"],
            (2160.0 / (buyin / 100.0)) / (stats["hands"] / 100.0),
        )


class TestTheRakebackSplit(unittest.TestCase):
    """The fall is priced twice, and the two prices have to be the same fall."""

    def setUp(self):
        self.config = _config(rakeback_pct=0.30, stakes=(
            Stake("100NL", 1.0, 7.46, 92.0, rake_bb100=8.0),
            Stake("200NL", 2.0, 4.32, 92.0, rake_bb100=6.0),
        ))
        self.allocation = mix.evaluate((4, 8), self.config)

    def _describe(self, equity, rakeback):
        equity = np.asarray(equity, dtype=float)
        rakeback = np.asarray(rakeback, dtype=float)
        hands = np.arange(len(equity)) * trials._HANDS_PER_STEP
        return trials._describe(self.config, self.allocation, equity, rakeback,
                                hands, False)

    def test_the_mean_splits_into_tables_plus_rakeback(self):
        # The split is taken out of the same sum `rates.total_winrate` builds,
        # so it must add back up exactly - a part that does not reconstitute the
        # whole would put every number on these slides out of step with the
        # main deck's EV line.
        self.assertGreater(self.allocation.rakeback_eur_per_100, 0.0)
        self.assertAlmostEqual(
            self.allocation.table_mean_eur_per_100
            + self.allocation.rakeback_eur_per_100,
            self.allocation.mean_eur_per_100,
        )

    def test_no_rakeback_leaves_the_mean_untouched(self):
        config = _config(rakeback_pct=0.0, stakes=self.config.stakes)
        allocation = mix.evaluate((4, 8), config)
        self.assertEqual(allocation.rakeback_eur_per_100, 0.0)
        self.assertAlmostEqual(allocation.table_mean_eur_per_100,
                               allocation.mean_eur_per_100)

    def test_the_fall_is_measured_on_the_cards_not_on_the_bankroll(self):
        # Totals rise 0 -> 60 -> 40 -> 100 while rakeback ramps 20 a step. The
        # cards therefore paid 0, 40, 0, 40: a fall of 40, not the 20 the total
        # dipped. Netting the rebate off first would report the smaller one.
        stats = self._describe([0, 60, 40, 100], [0, 20, 40, 60])
        self.assertAlmostEqual(stats["drawdown_eur"], 40.0)
        self.assertEqual(stats["drawdown_from"], 1 * trials._HANDS_PER_STEP)
        self.assertEqual(stats["drawdown_to"], 2 * trials._HANDS_PER_STEP)

    def test_the_net_fall_is_the_drop_in_the_bankroll_over_the_same_span(self):
        # Not a second drawdown measurement - the same two hand counts, read off
        # the other line. That identity is what lets the slide say "a 40 buy-in
        # fall at the tables was a 20 buy-in fall in the roll" and have both
        # numbers describe one event.
        equity = [0, 60, 40, 100]
        stats = self._describe(equity, [0, 20, 40, 60])
        self.assertAlmostEqual(stats["drawdown_rakeback_eur"], 20.0)
        self.assertAlmostEqual(stats["drawdown_net_eur"], 20.0)
        self.assertAlmostEqual(stats["drawdown_net_eur"],
                               equity[1] - equity[2])

    def test_the_money_rows_sum_to_the_total(self):
        stats = self._describe([0, 60, 40, 100], [0, 20, 40, 60])
        self.assertAlmostEqual(stats["won_table_eur"] + stats["won_rakeback_eur"],
                               stats["won_eur"])
        self.assertAlmostEqual(stats["won_table_buyins"]
                               + stats["won_rakeback_buyins"], stats["won_buyins"])
        self.assertAlmostEqual(stats["bb100_table"] + stats["bb100_rakeback"],
                               stats["bb100"])

    def test_the_bankroll_recovers_first_when_rakeback_keeps_arriving(self):
        # The cards pay 0, 40, 0, 10, 40 against a rebate ramping 20 a step, so
        # the total is 0, 60, 40, 70, 120. The cards take two steps to get
        # back to their 40 from the trough; the roll is back to its 60 after
        # one, carried the rest of the way by the rebate.
        stats = self._describe([0, 60, 40, 70, 120], [0, 20, 40, 60, 80])
        self.assertEqual(stats["recovery_hands"], 2 * trials._HANDS_PER_STEP)
        self.assertEqual(stats["recovery_net_hands"], 1 * trials._HANDS_PER_STEP)

    def test_a_ruined_lifetime_stops_earning_rakeback(self):
        # No hands, no rebate. If the ramp carried on past the barrier the table
        # line would keep falling underneath a flat total, which never happened.
        config = _config(bankroll_eur=1_500.0, rakeback_pct=0.30,
                         stakes=self.config.stakes)
        allocation = mix.evaluate((0, 12), config)
        ruined = [t for t in trials.run_trials(config, allocation, "optimal",
                                               count=40) if t.stats["ruined"]]
        self.assertTrue(ruined, "expected at least one busted lifetime")
        for trial in ruined:
            at_barrier = np.nonzero(trial.equity_eur <= -config.bankroll_eur)[0][0]
            self.assertAlmostEqual(trial.rakeback_eur[-1],
                                   trial.rakeback_eur[at_barrier])
            # And the split still holds where it matters - flat, not diverging.
            self.assertAlmostEqual(trial.table_eur[-1], trial.table_eur[at_barrier])

    def test_the_axis_covers_the_table_line_not_just_the_total(self):
        # The table line runs BELOW the total by the rakeback earned so far, so
        # scaling off the total alone would clip it off the bottom.
        drawn = trials.run_trials(self.config, self.allocation, "optimal", count=6)
        low, high = trials.trial_scales(self.config, drawn)
        for trial in drawn:
            self.assertLessEqual(low, trial.table_eur.min())
            self.assertGreaterEqual(high, trial.equity_eur.max())


class TestTrialPaths(unittest.TestCase):
    def setUp(self):
        self.config = _config()
        self.allocation = mix.evaluate((4, 8), self.config)

    def test_the_split_is_exact_at_every_step(self):
        config = _config(rakeback_pct=0.30, stakes=(
            Stake("100NL", 1.0, 7.46, 92.0, rake_bb100=8.0),
            Stake("200NL", 2.0, 4.32, 92.0, rake_bb100=6.0),
        ))
        allocation = mix.evaluate((4, 8), config)
        for trial in trials.run_trials(config, allocation, "optimal", count=3):
            np.testing.assert_allclose(
                trial.table_eur + trial.rakeback_eur, trial.equity_eur)
            # And the rebate is a straight ramp, because it is paid on hands
            # dealt: no variance, so no bends.
            steps = np.diff(trial.rakeback_eur)
            np.testing.assert_allclose(steps, steps[0])

    def test_full_resolution_is_kept(self):
        drawn = trials.run_trials(self.config, self.allocation, "optimal", count=3)
        self.assertEqual(len(drawn), 3)
        for trial in drawn:
            # One point per 100-hand block, plus the start. The whole reason
            # this module does not reuse `sim`'s checkpoints.
            self.assertEqual(
                len(trial.equity_eur),
                self.config.timescale_hands // trials._HANDS_PER_STEP + 1,
            )
            self.assertEqual(trial.equity_eur[0], 0.0)

    def test_the_same_seed_reproduces_the_same_lifetimes(self):
        # A comment pinned to "trial 7" has to still describe trial 7 after a
        # rebuild.
        first = trials.run_trials(self.config, self.allocation, "optimal", count=4)
        again = trials.run_trials(self.config, self.allocation, "optimal", count=4)
        for a, b in zip(first, again):
            np.testing.assert_allclose(a.equity_eur, b.equity_eur)

    def test_ruin_is_absorbing(self):
        # A bankroll small enough that some of the paths bust. Once busted, a
        # lifetime must sit on the barrier rather than trading its way back.
        config = _config(bankroll_eur=1_500.0)
        allocation = mix.evaluate((0, 12), config)
        for trial in trials.run_trials(config, allocation, "optimal", count=40):
            if trial.stats["ruined"]:
                self.assertAlmostEqual(trial.equity_eur[-1], -config.bankroll_eur)
                self.assertGreaterEqual(trial.equity_eur.min(), -config.bankroll_eur)

    def test_every_slide_gets_an_identically_sized_frame(self):
        # The section is built to be FLICKED THROUGH, which only works if the
        # plotting area is in the same place on every slide. It is not
        # automatic: the label naming the deepest fall is about half the width
        # of the axes and is not clipped, so wherever it overhangs the frame it
        # enters the axes' tight bounding box and `tight_layout` shrinks the
        # plot to fit it - by a different amount per lifetime, since the trough
        # lands somewhere different every time. `_annotate_fall` keeps it out of
        # the layout; this is what says so.
        import matplotlib.pyplot as plt

        drawn = trials.run_trials(self.config, self.allocation, "optimal", count=8)
        ylim = trials.trial_scales(self.config, drawn)
        boxes = set()
        for trial in drawn:
            figure = trials.trial_figure(trial, self.config, ylim=ylim)
            box = figure.axes[0].get_position()
            boxes.add((round(box.x0, 6), round(box.y0, 6),
                       round(box.x1, 6), round(box.y1, 6)))
            plt.close(figure)
        self.assertEqual(len(boxes), 1, f"frames differ between slides: {boxes}")

    def test_the_axis_covers_every_path_and_clears_the_ruin_barrier(self):
        drawn = trials.run_trials(self.config, self.allocation, "optimal", count=6)
        low, high = trials.trial_scales(self.config, drawn)
        self.assertLessEqual(low, -self.config.bankroll_eur)
        for trial in drawn:
            self.assertLessEqual(low, trial.equity_eur.min())
            self.assertGreaterEqual(high, trial.equity_eur.max())


class TestSimsDeck(unittest.TestCase):
    """The deck itself - structural, like `test_deck`."""

    @classmethod
    def setUpClass(cls):
        from pptx import Presentation

        cls.config = _config(timescale_hands=100_000)
        allocations = mix.all_allocations(cls.config)
        cls.best = mix.best_allocation(allocations, cls.config)
        cls.current = mix.evaluate((6, 6), cls.config)
        cls.path = Path(tempfile.mkdtemp()) / "sims.pptx"
        trials.write(cls.path, cls.config, cls.best, cls.current)
        cls.prs = Presentation(cls.path)

    def test_a_divider_a_summary_and_a_slide_per_trial_for_each_mix(self):
        self.assertEqual(len(list(self.prs.slides)), 2 * (2 + trials.TRIALS))

    def test_every_trial_slide_carries_its_numbers(self):
        tables = [
            shape.table
            for slide in self.prs.slides
            for shape in slide.shapes
            if shape.has_table
        ]
        # Two summaries plus one corner table per trial slide.
        self.assertEqual(len(tables), 2 + 2 * trials.TRIALS)

    def test_the_overlaid_table_stays_in_the_corner_it_is_meant_to(self):
        # It sits ON the chart by design, in the TOP-LEFT: the corner nothing is
        # ever drawn in, because every lifetime starts at zero and the axis is
        # scaled to where the best of twenty finishes. Anywhere lower and it
        # covers either the path or the label naming the deepest fall.
        for slide in self.prs.slides:
            picture = next((s for s in slide.shapes if s.shape_type == 13), None)
            tables = [s for s in slide.shapes if s.has_table]
            if picture is None or not tables:
                continue  # the divider and the summary slides
            for shape in tables:
                self.assertGreaterEqual(shape.left, picture.left)
                self.assertGreaterEqual(shape.top, picture.top)
                # Inside the top-left quadrant, both ways. The width bound is
                # what keeps the hottest lifetime's own peak out from under it;
                # it is measured on the PICTURE, whose left twentieth is the
                # axis label and tick gutter, so it is a shade looser than the
                # 44% of the plotting area `_TABLE_WIDTHS` is sized against.
                self.assertLessEqual(shape.left + shape.width,
                                     picture.left + 0.47 * picture.width)
                self.assertLessEqual(shape.top + shape.height,
                                     picture.top + 0.45 * picture.height)

    def test_the_summary_grid_fits_the_slide(self):
        # Twelve columns is as many as the width takes, and the falls needed
        # four of them - so this is worth an assertion rather than an eyeball.
        from shotopt import pptx_common as pc

        width = sum(w for _, w, _ in trials._SUMMARY_COLS)
        self.assertLessEqual(pc.Inches(width), pc.CONTENT_WIDTH)

    def test_no_table_runs_off_the_slide(self):
        from shotopt import pptx_common as pc

        for slide in self.prs.slides:
            for shape in slide.shapes:
                if shape.has_table:
                    self.assertLessEqual(shape.left + shape.width, pc.SLIDE_WIDTH)
                    self.assertLessEqual(shape.top + shape.height, pc.SLIDE_HEIGHT)

    def test_it_builds_with_only_one_mix(self):
        from pptx import Presentation

        path = Path(tempfile.mkdtemp()) / "sims.pptx"
        trials.write(path, self.config, self.best, None)
        prs = Presentation(path)
        self.assertEqual(len(list(prs.slides)), 2 + trials.TRIALS)


class TestDeckWiring(unittest.TestCase):
    """`sims.pptx` is written from inside `deck.build`, on the same two mixes."""

    def test_the_main_deck_writes_the_companion_when_asked(self):
        from shotopt import deck

        config = _config(timescale_hands=100_000)
        directory = Path(tempfile.mkdtemp())
        sims_path = directory / "sims.pptx"
        main = deck.build(config, directory, sims_path=sims_path)
        self.assertTrue(main.exists())
        self.assertTrue(sims_path.exists())

    def test_without_the_path_only_the_main_deck_is_written(self):
        from shotopt import deck

        config = _config(timescale_hands=100_000)
        directory = Path(tempfile.mkdtemp())
        deck.build(config, directory)
        self.assertEqual([p.name for p in directory.glob("*.pptx")],
                         ["stake_optimisation.pptx"])


if __name__ == "__main__":
    unittest.main()
